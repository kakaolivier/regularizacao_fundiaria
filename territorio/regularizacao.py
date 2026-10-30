"""Typed land records and conservative import of BANCO_REG.

The dictionary is metadata, not executable code. Excel formulas/VBA are never run.
CNJ groups a case; only the immutable land-record key identifies an update.
"""
import hashlib
import json
import math
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .services import Problem, cnj, text
from .store import now

SCHEMA = json.loads(Path(__file__).with_name('regularizacao_schema.json').read_text(encoding='utf-8'))
FIELDS = {f['column']: f for f in SCHEMA['fields']}
GROUPS = {
    'CADASTRO': 'Cadastro', 'ATE_IMISSAO_DE_POSSE': 'Posse e obras',
    'FASE_PERICIAL': 'Perícia', 'ACORDO': 'Acordos', 'REGISTRO': 'Registro',
    'FASE_FINAL': 'Fase final', 'ATUALIZACAO_MVGA': 'Acompanhamento MVGA', 'GIS': 'GIS',
}
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
MISSING = {'', '-', 'n/a', 'na', 'nao se aplica'}


def normalize(value):
    return ''.join(c for c in unicodedata.normalize('NFD', str(value).strip().casefold()) if unicodedata.category(c) != 'Mn')

def clean(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def meaningful(value):
    return normalize(value) not in MISSING

def digest(value):
    return hashlib.sha256(value.encode('utf-8')).hexdigest()

def iso_date(value):
    if not meaningful(value):
        return None
    try:
        return date.fromisoformat(value) if re.fullmatch(r'\d{4}-\d{2}-\d{2}', value) else None
    except ValueError:
        return None

def computed(values):
    start, end = iso_date(values.get('U', '')), iso_date(values.get('V', ''))
    return {
        'W': str((end - start).days) if start and end else '',
        'AP': 'LIBERADO PARA OBRA' if iso_date(values.get('AO', '')) else 'AGUARDANDO LIBERAÇÃO',
    }

def problems(values):
    warnings = []
    for col, field in FIELDS.items():
        value = values.get(col, '')
        if not meaningful(value):
            if field['required'] and field['type'] != 'auto':
                warnings.append({'column': col, 'message': f"{field['label']}: não informado / não aplicável; conferir no contexto do caso."})
            continue
        if field['type'] == 'choice' and normalize(value) not in {normalize(v) for v in field['options']}:
            warnings.append({'column': col, 'message': f"{field['label']}: valor fora da lista do dicionário; original preservado."})
        if field['type'] == 'date' and not iso_date(value):
            warnings.append({'column': col, 'message': f"{field['label']}: texto sem data exata; original preservado."})
    derived = computed(values)
    for col in ('W', 'AP'):
        if meaningful(values.get(col, '')) and normalize(values[col]) != normalize(derived[col]):
            warnings.append({'column': col, 'message': f"{FIELDS[col]['label']}: valor de origem difere do cálculo pelas datas."})
    if derived['W'] and int(derived['W']) < 0:
        warnings.append({'column': 'W', 'message': 'Petição anterior à solicitação de ajuizamento; conferir as datas.'})
    return warnings

def _cell_value(cell, workbook, sheet):
    value = cell.find(NS + 'v')
    raw = value.text if value is not None else None
    kind = cell.get('t')
    if kind == 's' and raw is not None:
        return sheet._shared_strings[int(raw)]
    if kind == 'inlineStr':
        return ''.join(n.text or '' for n in cell.iter(NS + 't'))
    if kind == 'e':
        return 'Erro de fórmula: ' + (raw or 'sem resultado')
    if raw is not None and kind not in ('str', 'd'):
        try:
            value = float(raw)
            if not math.isfinite(value):
                return raw
            if int(cell.get('s', 0)) in workbook._date_formats:
                return from_excel(value, workbook.epoch)
            return int(value) if value.is_integer() else value
        except (ValueError, OverflowError):
            pass
    return raw

def _typed(col, value, epoch):
    if FIELDS[col]['type'] == 'date':
        if isinstance(value, (int, float)) and 1 <= value <= 2958465:
            try:
                return clean(from_excel(value, epoch))
            except (ValueError, OverflowError):
                pass
        if isinstance(value, str) and re.fullmatch(r'\d{2}/\d{2}/\d{4}', value.strip()):
            try:
                return datetime.strptime(value.strip(), '%d/%m/%Y').date().isoformat()
            except ValueError:
                pass
    return clean(value)

def read_control(path, filename=None):
    """Stream real cells, ignoring Excel's million formatted empty rows."""
    path = Path(path)
    try:
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 300 * 1024 * 1024:
                raise Problem('Planilha expandida excede 300 MB. Remova linhas vazias formatadas e tente novamente.')
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, KeyError, ValueError, ET.ParseError, OSError) as exc:
        raise Problem('Não foi possível ler o XLSX. Selecione uma pasta de trabalho válida.') from exc
    try:
        if 'BANCO_REG' not in workbook.sheetnames:
            raise Problem('Esta importação espera a aba BANCO_REG de CONTROLE_REGULARIZACAO.xlsx. O dicionário já está incorporado ao site.')
        sheet = workbook['BANCO_REG']
        expected = {re.sub(r'[^a-z0-9]', '', normalize(f['title'])): c for c, f in FIELDS.items()}
        mapping, rows, physical = {}, [], 0
        with workbook._archive.open(sheet._worksheet_path) as stream:
            parent = None
            for event, element in ET.iterparse(stream, events=('start', 'end')):
                if event == 'start' and element.tag == NS + 'sheetData':
                    parent = element
                if event != 'end' or element.tag != NS + 'row':
                    continue
                physical += 1
                line = int(element.get('r'))
                raw, formulas = {}, []
                for cell in element:
                    col = re.sub(r'\d+', '', cell.get('r', ''))
                    value = _cell_value(cell, workbook, sheet)
                    if value is not None:
                        raw[col] = value
                    if cell.find(NS + 'f') is not None and value is None:
                        formulas.append(col)
                if line == 3:
                    for col, value in raw.items():
                        canonical = expected.get(re.sub(r'[^a-z0-9]', '', normalize(value)))
                        if not canonical and clean(value):
                            raise Problem('Cabeçalho não previsto no dicionário: ' + clean(value)[:150] + '. Confira o layout antes de importar.')
                        if canonical:
                            if canonical in mapping.values():
                                raise Problem('Há cabeçalhos repetidos na linha 3 de BANCO_REG.')
                            mapping[col] = canonical
                    if set(mapping.values()) != set(FIELDS):
                        missing = ', '.join(FIELDS[c]['label'] for c in FIELDS if c not in mapping.values())
                        raise Problem('Layout diferente do dicionário. Campos ausentes na linha 3: ' + missing)
                elif line > 3:
                    if not mapping:
                        raise Problem('Cabeçalho de BANCO_REG não encontrado na linha 3.')
                    values = {mapping[c]: _typed(mapping[c], v, workbook.epoch) for c, v in raw.items() if c in mapping}
                    if any(meaningful(values.get(c, '')) for c in ('E', 'F', 'G', 'H', 'L', 'B')):
                        if len(rows) >= 10000:
                            raise Problem('Limite de 10.000 registros por importação.')
                        key = '|'.join(values.get(c, '') for c in ('L', 'F', 'G'))
                        supplied = values.get('C', '')
                        errors = []
                        if supplied and [part.strip() for part in supplied.split('|')] != [values.get(c, '') for c in ('L', 'F', 'G')]:
                            errors.append('Chave de origem não corresponde a processo + imóvel + subtrecho. Confira no Excel.')
                        if not all(meaningful(values.get(c, '')) for c in ('E', 'F', 'G')):
                            errors.append('Município, imóvel e subtrecho são necessários para identificar este registro.')
                        if any(len(v) > 20000 for v in values.values()):
                            errors.append('Campo excede 20.000 caracteres.')
                        values['C'] = supplied or key
                        warnings = problems(values)
                        for col in formulas:
                            if mapping.get(col) in FIELDS and FIELDS[mapping[col]]['type'] != 'auto':
                                warnings.append({'column': mapping[col], 'message': 'Fórmula sem resultado salvo. Recalcule no Excel; nenhuma fórmula foi executada.'})
                        rows.append({'row': line, 'key': values['C'], 'values': values, 'warnings': warnings, 'errors': errors,
                                     'rawFields': [{'column': c, 'label': FIELDS[mapping[c]]['title'], 'value': clean(v)} for c, v in raw.items() if c in mapping]})
                element.clear()
                if parent is not None:
                    parent.clear()
        if not mapping or not rows:
            raise Problem('Nenhum registro foi encontrado em BANCO_REG.')
        counts = Counter(r['key'] for r in rows)
        for row in rows:
            if counts[row['key']] > 1:
                row['errors'].append('Chave repetida na planilha. Diferencie os subtrechos antes de importar essas linhas.')
        return {'filename': Path(filename or path.name).name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'rows': rows, 'ignoredRows': max(0, physical - len(rows) - 2), 'duplicateGroups': sum(v > 1 for v in counts.values()), 'createdAt': now()}
    except (KeyError, ValueError, IndexError, OverflowError, ET.ParseError) as exc:
        raise Problem('Não foi possível interpretar os dados do XLSX. Confira o layout de BANCO_REG.') from exc
    finally:
        workbook.close()

def fingerprint(store):
    return digest(json.dumps([store.all('records'), store.all('cases'), [m['id'] for m in store.all('meta')]], sort_keys=True, ensure_ascii=False))

def plan_import(store, bundle):
    records = {r['key']: r for r in store.all('records')}
    cases = defaultdict(list)
    for case in store.all('cases'):
        number = cnj(case.get('number', ''))
        if number:
            cases[number].append(case['id'])
    imported = store.get('meta', 'control-' + bundle['sha256'])
    items = []
    for row in bundle['rows']:
        old = records.get(row['key'])
        errors = list(row['errors'])
        changes = []
        if old:
            for col, value in row['values'].items():
                previous = old['values'].get(col, '')
                if not value or value == previous or FIELDS[col]['type'] == 'auto':
                    continue
                if not FIELDS[col]['overwrite'] and meaningful(previous):
                    errors.append(f"{FIELDS[col]['label']}: valor já cadastrado é protegido contra substituição.")
                else:
                    changes.append({'column': col, 'label': FIELDS[col]['label'], 'before': previous, 'after': value})
        else:
            number = cnj(row['values'].get('L', ''))
            if number and len(cases[number]) > 1:
                errors.append('Há mais de um processo existente com este CNJ. Resolva o vínculo antes da importação.')
        status = 'blocked' if errors else 'unchanged' if imported or (old and not changes) else 'updated' if old else 'new'
        items.append({**row, 'status': status, 'errors': errors, 'changes': changes, 'previousVersion': old['version'] if old else None,
                      'caseId': old['caseId'] if old else next(iter(cases.get(cnj(row['values'].get('L', '')), [])), None)})
    return {'items': items, 'counts': dict(Counter(i['status'] for i in items)), 'fingerprint': fingerprint(store),
            'alreadyImported': bool(imported), 'warnings': sum(bool(i['warnings']) for i in items)}

def _case_for(store, item):
    if item['caseId']:
        return item['caseId']
    values = item['values']
    number = cnj(values.get('L', ''))
    # Same CNJ can share a dossier; its separate keyed records never merge.
    matching = [c for c in store.all('cases') if number and cnj(c.get('number', '')) == number]
    if matching:
        return matching[0]['id']
    identifier = 'cnj-' + number if number else 'reg-' + digest(item['key'])[:24]
    kind = normalize(values.get('J', ''))
    method = 'Extrajudicial' if kind == 'extrajudicial' else 'Judicial' if 'judicial' in kind else 'Administrativo' if kind == 'administrativo' else 'A definir'
    case = {'id': identifier, 'title': values.get('F', 'Imóvel não informado'), 'number': values.get('L', ''),
            'city': values.get('E', ''), 'owner': values.get('H', ''), 'section': values.get('G', ''), 'method': method,
            'stage': values.get('N') or 'Etapa não informada', 'category': 'Em andamento', 'expertStage': '',
            'hasExpert': False, 'hasAward': False, 'hasRegistry': False, 'sourceCount': 0, 'sourceIds': [],
            'needsReconciliation': not bool(number), 'origins': ['controle'], 'createdAt': now(), 'regularizacaoManaged': True}
    store.put('cases', case)
    return identifier

def _source(item, bundle, case_id, record_id, import_id):
    values = item['values']
    return {'id': import_id + '-' + str(item['row']), 'recordId': record_id, 'workbook': 'controle', 'filename': bundle['filename'],
            'sheet': 'BANCO_REG', 'row': item['row'], 'kind': 'controle', 'title': values.get('F', ''), 'city': values.get('E', ''),
            'section': values.get('G', ''), 'owner': values.get('H', ''), 'process': values.get('L', ''), 'stage': values.get('N', ''),
            'notes': values.get('AU', ''), 'registry': values.get('S') or values.get('R', ''), 'gisId': values.get('B', ''),
            'area': values.get('Q') or values.get('P', ''), 'method': values.get('J', ''), 'release': values.get('AP', ''),
            'fields': item['rawFields'], 'financial': {}, 'events': [], 'importedAt': now(), 'sha256': bundle['sha256']}

def apply_import(store, bundle, expected, author):
    plan = plan_import(store, bundle)
    if expected != plan['fingerprint']:
        raise Problem('A base mudou após a prévia. Gere uma nova prévia antes de confirmar.', 409)
    if plan['alreadyImported']:
        raise Problem('Este arquivo já foi importado. Uma repetição não substitui alterações feitas no site.', 409)
    if not any(i['status'] in ('new', 'updated') for i in plan['items']):
        raise Problem('Não há alterações elegíveis para aplicar.')
    author = text({'author': author}, 'author')
    backup = store.directory / 'backups'
    backup.mkdir(exist_ok=True)
    import_id = 'import-' + uuid4().hex
    backup_name = import_id + '.sqlite'
    with sqlite3.connect(backup / backup_name) as copy:
        store.db.backup(copy)
    with store.transaction():
        if expected != fingerprint(store):
            raise Problem('A base mudou durante a confirmação. Gere uma nova prévia.', 409)
        for item in plan['items']:
            if item['status'] not in ('new', 'updated'):
                continue
            identifier = 'record-' + digest(item['key'])
            old = store.get('records', identifier)
            case_id = old['caseId'] if old else _case_for(store, item)
            values = dict(old['values']) if old else dict(item['values'])
            if old:
                values.update({change['column']: change['after'] for change in item['changes']})
            source = _source(item, bundle, case_id, identifier, import_id)
            record = {'id': identifier, 'key': item['key'], 'caseId': case_id, 'values': values,
                      'version': old['version'] + 1 if old else 1, 'updatedAt': now(),
                      'sourceIds': (old['sourceIds'] if old else []) + [source['id']], 'lastRow': item['row']}
            store.put('records', record, case_id)
            store.put('sources', source, case_id)
            case = store.get('cases', case_id)
            source_ids = case.get('sourceIds', []) + [source['id']]
            store.put('cases', {**case, 'sourceIds': source_ids, 'sourceCount': len(source_ids),
                                'origins': list(dict.fromkeys(case.get('origins', []) + ['controle']))})
            store.event(case_id, kind='activity', author=author, title='Imóvel atualizado pela planilha' if old else 'Imóvel / trecho importado',
                        body=f"{values.get('F', '')} · {values.get('G', '')} · versão {record['version']}",
                        recordId=identifier, changes=item['changes'], source=f"{bundle['filename']} · BANCO_REG · linha {item['row']}")
            if not old and meaningful(values.get('AU', '')):
                store.event(case_id, kind='source', author='Planilha de origem', date=None, title='Situação informada no controle',
                            body=values['AU'], recordId=identifier, source=f"BANCO_REG · linha {item['row']}")
        report = {'id': 'control-' + bundle['sha256'], 'filename': bundle['filename'], 'sha256': bundle['sha256'],
                  'importedAt': now(), 'counts': plan['counts'], 'backup': backup_name, 'author': author,
                  'blocked': [{'row': i['row'], 'key': i['key'], 'errors': i['errors']} for i in plan['items'] if i['status'] == 'blocked']}
        store.put('meta', report)
    return report

def record_view(record):
    values = record['values']
    return {**record, 'warnings': problems(values), 'computed': computed(values),
            'title': values.get('F', ''), 'section': values.get('G', ''), 'city': values.get('E', ''),
            'stage': values.get('N') or 'Etapa não informada', 'expertStage': values.get('AE', ''),
            'owner': values.get('H', ''), 'number': values.get('L', ''), 'assignee': values.get('BD', '')}

def update_record(store, identifier, data):
    author = text(data, 'author')
    reason = text(data, 'reason', 1000)
    group = text(data, 'group')
    if group not in GROUPS:
        raise Problem('Área de trabalho inválida.')
    with store.transaction():
        record = store.get('records', identifier)
        if not record:
            raise Problem('Imóvel / trecho não encontrado.', 404)
        if str(record['version']) != data.get('version'):
            raise Problem('Este registro foi alterado. Reabra a ficha antes de salvar.', 409)
        values = dict(record['values'])
        changes = []
        for col, field in FIELDS.items():
            if col not in data:
                continue
            if field['group'] != group:
                raise Problem('O campo não pertence à área de trabalho selecionada.')
            value = text(data, col, 20000, False)
            previous = values.get(col, '')
            if not value or value == previous:
                continue
            if field['type'] == 'auto' or (not field['overwrite'] and meaningful(previous)):
                raise Problem(f"{field['label']}: campo protegido contra edição.")
            if field['type'] == 'choice':
                options = {normalize(v): v for v in field['options']}
                if normalize(value) not in options:
                    raise Problem(f"{field['label']}: selecione um valor da lista.")
                value = options[normalize(value)]
            elif field['type'] == 'date' and not iso_date(value):
                raise Problem(f"{field['label']}: informe uma data completa válida.")
            elif field['type'] == 'money':
                if not re.fullmatch(r'\d+(?:[.,]\d{1,2})?', value):
                    raise Problem(f"{field['label']}: use um valor positivo sem separador de milhar.")
                value = value.replace(',', '.')
            values[col] = value
            changes.append({'column': col, 'label': field['label'], 'before': previous, 'after': value})
        if not changes:
            raise Problem('Nenhuma alteração preenchida. Campos vazios preservam o valor existente.')
        # The original key remains immutable even if a judicial number is formally corrected.
        updated = {**record, 'values': values, 'version': record['version'] + 1, 'updatedAt': now()}
        store.put('records', updated, record['caseId'])
        store.event(record['caseId'], kind='activity', author=author, title='Atualização de ' + GROUPS[group],
                    body=f"{values.get('F', '')} · {values.get('G', '')}\nMotivo / referência: {reason}", recordId=identifier, changes=changes)
        return updated
