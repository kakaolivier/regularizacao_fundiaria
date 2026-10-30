"""Generate fictional files for docs/MANUAL-DE-TESTES.md without touching a database."""
import argparse
import csv
import json
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
NUMBER = '9999999-00.2026.8.09.9999'  # Illustrative identifier, never a verified legal case.
TITLE = 'TESTE - EEE Jardim Modelo'


def format_sheet(sheet, header_row):
    sheet.freeze_panes = f'A{header_row + 1}'
    sheet.auto_filter.ref = f'A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}'
    for cell in sheet[header_row]:
        cell.fill = PatternFill('solid', fgColor='204E3C')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.column_dimensions[cell.column_letter].width = 26
    sheet.row_dimensions[header_row].height = 48
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = 'DD/MM/YYYY'


def write_control(path, rows, fields):
    book = Workbook()
    sheet = book.active
    sheet.title = 'BANCO_REG'
    sheet['A1'] = 'DADOS FICTÍCIOS — EXCLUSIVAMENTE PARA TESTES DO SOFTWARE'
    for field in fields:
        sheet[field['column'] + '2'] = field['group']
        sheet[field['column'] + '3'] = field['title']
    for line, row in enumerate(rows, 4):
        row = {**row, 'C': '|'.join(str(row.get(c) or '') for c in ('L', 'F', 'G'))}
        for column, value in row.items():
            sheet[f'{column}{line}'] = value
    format_sheet(sheet, 3)
    book.save(path)
    book.close()


def write_legacy(directory):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Processos ajuizados'
    sheet['A1'] = 'DADOS FICTÍCIOS — NÃO UTILIZAR EM PROCESSOS REAIS'
    headers = {'C': 'ID GIS', 'E': 'MUNICÍPIO', 'F': 'DENOMINAÇÃO', 'G': 'SUBTRECHO',
               'H': 'PROPRIETÁRIO', 'I': 'PROCESSO', 'J': 'ETAPA', 'N': 'ÁREA', 'P': 'MATRÍCULA',
               'S': 'SOLICITAÇÃO', 'T': 'PETIÇÃO', 'V': 'INDENIZAÇÃO DEPOSITADA', 'W': 'PROPOSTA',
               'X': 'LAUDO PERICIAL', 'AB': 'SENTENÇA', 'AD': 'ACORDO', 'AN': 'SITUAÇÃO'}
    values = {'C': 'GIS-TESTE-001', 'E': 'Aparecida de Goiânia', 'F': TITLE, 'G': 'Trecho 1',
              'H': 'Pessoa fictícia A', 'I': NUMBER, 'J': 'REALIZAÇÃO DE PERÍCIA', 'N': 250,
              'P': 'MAT-TESTE-001', 'S': datetime(2026, 10, 1), 'T': datetime(2026, 10, 6),
              'V': 10000, 'W': 15000, 'X': 15000, 'AB': 14000, 'AD': 14000,
              'AN': 'Referências financeiras fictícias. Não representam um processo real.'}
    for col, label in headers.items():
        sheet[col + '5'] = label
    for col, value in values.items():
        sheet[col + '6'] = value
    format_sheet(sheet, 5)
    book.save(directory / '04_ACOMPANHAMENTO_TESTE.xlsx')
    book.close()

    book = Workbook()
    sheet = book.active
    sheet.title = 'LAUDOS APRESENTADOS'
    sheet['A1'] = 'DADOS FICTÍCIOS — NÃO UTILIZAR EM PROCESSOS REAIS'
    headers = {'B': 'MUNICÍPIO', 'C': 'IMÓVEL', 'D': 'SUBTRECHO', 'E': 'PROCESSO', 'F': 'PARTE',
               'G': 'DEPÓSITO', 'H': 'DEPÓSITO ATUALIZADO', 'J': 'PRIMEIRO LAUDO', 'K': 'SEGUNDO LAUDO',
               'L': 'LAUDO ATUALIZADO', 'M': 'HONORÁRIOS SUCUMBENCIAIS', 'O': 'A DESEMBOLSAR', 'S': 'FASE', 'T': 'NOTAS'}
    values = {'B': 'Aparecida de Goiânia', 'C': TITLE, 'D': 'Trecho 1', 'E': NUMBER, 'F': 'Pessoa fictícia A',
              'G': 10000, 'H': 10500, 'J': 15000, 'K': 16500, 'L': 16000, 'M': 1000, 'O': 7000,
              'S': 'LAUDO APRESENTADO', 'T': 'Estimativas fictícias, sem cálculo jurídico/financeiro validado.'}
    for col, label in headers.items():
        sheet[col + '11'] = label
    for col, value in values.items():
        sheet[col + '12'] = value
    format_sheet(sheet, 11)
    book.save(directory / '05_PERICIAS_TESTE.xlsx')
    book.close()


def generate(directory):
    directory = Path(directory)
    if directory.exists() and any(directory.iterdir()):
        raise SystemExit('A pasta de saída já contém arquivos. Reutilize-os ou escolha outra pasta com --output. Nenhum arquivo foi substituído.')
    directory.mkdir(parents=True, exist_ok=True)
    fields = json.loads((ROOT / 'territorio/regularizacao_schema.json').read_text(encoding='utf-8'))['fields']
    base = {'D': 1, 'E': 'APARECIDA DE GOIÂNIA', 'F': TITLE, 'G': 'Trecho 1', 'H': 'Pessoa fictícia A',
            'J': 'JUDICIAL', 'K': 'PROTOCOLO-TESTE-001', 'L': NUMBER, 'N': 'REALIZAÇÃO DE PERÍCIA',
            'P': 250, 'R': 'MAT-TESTE-001', 'T': 'DUP FICTÍCIA PARA TESTES',
            'U': datetime(2026, 10, 1), 'V': datetime(2026, 10, 6), 'W': 5, 'X': 10000,
            'Z': 'NÃO', 'AE': 'PERITO NOMEADO', 'AN': 'IMISSÃO PROVISÓRIA AINDA NÃO APRECIADA',
            'AP': 'AGUARDANDO LIBERAÇÃO', 'AU': 'Preparar conferência fictícia do cadastro.',
            'BD': 'Equipe técnica fictícia'}
    second = {**deepcopy(base), 'D': 2, 'G': 'Trecho 2', 'H': 'Pessoa fictícia B', 'P': 150, 'R': 'MAT-TESTE-002', 'X': 5000}
    write_control(directory / '01_CONTROLE_TESTE.xlsx', [base, second], fields)
    # Leave editable fields blank in a revision, so earlier exercises are preserved.
    revision = {col: base[col] for col in ('D', 'E', 'F', 'G', 'J', 'L')}
    revision.update(AU='Providência de revisão fictícia: conferir matrícula.', BD='Equipe de revisão fictícia')
    write_control(directory / '02_REVISAO_TESTE.xlsx', [revision, second], fields)
    duplicate = {**base, 'F': 'TESTE - Chave duplicada', 'G': 'ÁREA'}
    incomplete = {**base, 'D': 3, 'F': '', 'G': 'Sem denominação'}
    approximate = {**base, 'D': 4, 'F': 'TESTE - Previsão aproximada', 'G': 'ÁREA', 'BC': 'OUT/2026'}
    write_control(directory / '03_PENDENCIAS_TESTE.xlsx', [duplicate, {**duplicate, 'D': 2, 'P': 180}, incomplete, approximate], fields)
    write_legacy(directory)
    feature = {'type': 'Feature', 'properties': {'name': 'ÁREA FICTÍCIA PARA TESTES'}, 'geometry': {'type': 'Polygon',
               'coordinates': [[[-49.25, -16.82], [-49.249, -16.82], [-49.249, -16.819], [-49.25, -16.819], [-49.25, -16.82]]]}}
    (directory / 'area_teste.geojson').write_text(json.dumps(feature, ensure_ascii=False, indent=2), encoding='utf-8')
    for version in (1, 2):
        (directory / f'documento_teste_v{version}.txt').write_text(f'DOCUMENTO FICTÍCIO — SEM VALIDADE JURÍDICA\nVersão {version}\nArquivo de exemplo para testar envio, download e versionamento.\n', encoding='utf-8')
    scenarios = [
        ('T01', 'Importar controle', '2 imóveis / trechos em 1 processo'),
        ('T02', 'Indenização depositada', 'Trecho 1: R$ 12.000,50; trecho 2: R$ 5.000,00'),
        ('T03', 'Perícia', 'Laudos de R$ 15.000,25 e R$ 16.500,00 preservados separadamente'),
        ('T04', 'Acordos', 'R$ 14.000,00; laudos anteriores preservados'),
        ('T05', 'Posse e obras', '5 dias calculados; indicador pela data do mandado atualizado'),
        ('T06', 'Registro e acompanhamento', 'Campos salvos com motivo e histórico'),
        ('T07', 'Andamento', 'Evento do processo salvo sem mudar etapa automaticamente'),
        ('T08', 'Documentos', 'Versões 1 e 2 disponíveis e downloads conferidos'),
        ('T09', 'Tarefas', 'Criar, concluir e reabrir'),
        ('T10', 'Mapa', 'Área salva, vinculada ao Trecho 1, editada e exportada'),
        ('T11', 'Persistência', 'Reabrir a mesma base conserva registros e anexos'),
        ('T12', 'Revisão', '1 atualizado / 1 sem alterações; valores editados preservados'),
        ('T13', 'Pendências', 'Prévia com 1 novo / 3 bloqueados; não confirmar este exercício'),
        ('T14', 'Regras de preenchimento', 'Vazio preserva; formato monetário inválido é rejeitado'),
        ('T15', 'Conflito entre abas', 'Segunda gravação desatualizada é recusada'),
        ('T16', 'Valores e referências', 'Duas fontes anteriores visíveis, sem soma ou edição direta'),
    ]
    with (directory / 'roteiro_resultados.csv').open('w', newline='', encoding='utf-8-sig') as file:
        writer = csv.writer(file, delimiter=';')
        writer.writerow(['Teste', 'Cenário', 'Resultado esperado', 'Resultado observado', 'Status', 'Observações'])
        for identifier, scenario, expected in scenarios:
            writer.writerow([identifier, scenario, expected, '', 'Não executado', ''])
    (directory / 'LEIA-ME.txt').write_text('MATERIAL INTEIRAMENTE FICTÍCIO\nSiga docs/MANUAL-DE-TESTES.md.\n01, 02 e 03: importar pela interface, na ordem do manual.\n04 e 05: exercício opcional das fontes anteriores em outra base.\nNão importar estes exemplos na base de trabalho real.\nEste gerador não cria nem altera bancos de dados.\n', encoding='utf-8')
    print(f'Material fictício criado em: {directory.resolve()}')
    print('Abra docs/MANUAL-DE-TESTES.md e use iniciar_testes.bat para trabalhar em uma base separada.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Criar exemplos fictícios para o manual de testes.')
    parser.add_argument('--output', type=Path, default=ROOT / 'material-teste')
    generate(parser.parse_args().output)
