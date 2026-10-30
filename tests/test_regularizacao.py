import io
import json
import sqlite3
import tempfile
import unittest
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from territorio import create_app
from territorio import regularizacao as reg
from territorio.services import Problem
from territorio.store import Store

CNJ = '1234567-89.2025.8.09.0001'


def workbook(path, rows=None):
    book = Workbook()
    sheet = book.active
    sheet.title = 'BANCO_REG'
    for col, field in reg.FIELDS.items():
        sheet[f'{col}3'] = field['title']
    rows = rows if rows is not None else [{'G': 'Trecho 1'}, {'G': 'Trecho 2', 'P': 80}]
    for line, changed in enumerate(rows, 4):
        values = {'D': line - 3, 'E': 'APARECIDA DE GOIÂNIA', 'F': 'Área fictícia', 'G': 'Trecho 1',
                  'H': 'Pessoa fictícia', 'J': 'JUDICIAL', 'L': CNJ, 'P': 100,
                  'U': 45000, 'V': 45010, 'Z': 'SIM', 'AE': 'PERITO NOMEADO', 'X': 0,
                  'AO': datetime(2025, 1, 10), 'BC': 'LIBERADO', 'R': 'N/A', 'BD': 'Equipe técnica', **changed}
        if 'C' not in values:
            values['C'] = '|'.join(str(values.get(c) or '') for c in ('L', 'F', 'G'))
        for col, value in values.items():
            sheet[f'{col}{line}'] = value
    # Formatting at Excel's last row must not create fictitious records.
    sheet['D1048576'].fill = PatternFill('solid', fgColor='FF0000')
    extra = book.create_sheet('FASE_PERICIAL')
    extra['A1'] = 'Illustrative VBA / instructions are not executed or imported as cases.'
    book.save(path)
    book.close()
    return path


class RegularizacaoTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.app = create_app(self.root / 'data', testing=True)
        self.client = self.app.test_client()
        self.store = Store(self.root / 'data')

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def bundle(self, rows=None, name='controle.xlsx'):
        return reg.read_control(workbook(self.root / name, rows))

    def apply(self, bundle):
        plan = reg.plan_import(self.store, bundle)
        return reg.apply_import(self.store, bundle, plan['fingerprint'], 'Equipe de teste')

    def test_parser_dates_missing_values_duplicates_and_cached_formulas(self):
        path = workbook(self.root / 'source.xlsx', [{'G': 'A'}, {'G': 'B'}, {'G': 'B', 'P': 50}, {'E': None, 'G': 'C'}])
        original = path.read_bytes()
        bundle = reg.read_control(path)
        self.assertEqual(len(bundle['rows']), 4)
        self.assertEqual(bundle['duplicateGroups'], 1)
        self.assertEqual(bundle['rows'][0]['values']['U'], '2023-03-15')
        self.assertEqual(bundle['rows'][0]['values']['AO'], '2025-01-10')
        self.assertEqual(bundle['rows'][0]['values']['BC'], 'LIBERADO')
        self.assertEqual(bundle['rows'][0]['values']['X'], '0')
        self.assertEqual(bundle['rows'][0]['values']['R'], 'N/A')
        self.assertEqual(reg.computed(bundle['rows'][0]['values'])['W'], '10')
        self.assertEqual(reg.plan_import(self.store, bundle)['counts'], {'new': 1, 'blocked': 3})
        self.assertEqual(path.read_bytes(), original)
        # An uncached key formula is computed from the three cells, never evaluated.
        formula_bundle = self.bundle([{'C': '=untrusted_formula()', 'G': 'Formula'}], 'formula.xlsx')
        self.assertEqual(formula_bundle['rows'][0]['key'], CNJ + '|Área fictícia|Formula')

    def test_same_process_keeps_separate_records_and_existing_work(self):
        case = self.client.post('/api/cases', json={'title': 'Dossiê existente', 'city': 'Goiânia', 'owner': 'Parte anterior', 'number': CNJ, 'method': 'Judicial'}).json
        self.client.post(f"/api/cases/{case['id']}/documents", data={'category': 'Outros', 'file': (io.BytesIO(b'Arquivo existente'), 'anexo.txt')})
        self.client.post(f"/api/cases/{case['id']}/tasks", json={'title': 'Trabalho existente', 'assignee': 'Equipe', 'due': '2026-12-31'})
        self.client.post(f"/api/cases/{case['id']}/events", json={'title': 'Histórico anterior', 'body': 'Preservar', 'author': 'Equipe', 'kind': 'activity'})
        report = self.apply(self.bundle())
        self.assertEqual(report['counts'], {'new': 2})
        self.assertEqual(len(self.store.all('cases')), 1)
        records = self.store.all('records')
        self.assertEqual(len(records), 2)
        self.assertEqual({r['caseId'] for r in records}, {case['id']})
        before_backup = self.root / 'data' / 'backups' / report['backup']
        with sqlite3.connect(before_backup) as backup:
            self.assertEqual(backup.execute('SELECT count(*) FROM records').fetchone()[0], 0)
            self.assertEqual(backup.execute('SELECT count(*) FROM documents').fetchone()[0], 1)
        detail = self.client.get(f"/api/cases/{case['id']}").json
        self.assertEqual(detail['title'], 'Dossiê existente')
        self.assertEqual(len(detail['documents']), 1)
        self.assertEqual(len(detail['tasks']), 1)
        self.assertEqual(len(detail['records']), 2)
        self.assertTrue(any(e['title'] == 'Histórico anterior' for e in detail['events']))

    def test_process_type_filters_and_protected_case_level_stage(self):
        self.apply(self.bundle([{'J': 'ADMINISTRATIVO', 'G': 'Administrativo', 'L': ''}, {'J': 'EXTRAJUDICIAL', 'G': 'Extrajudicial', 'L': ''}]))
        cases = self.client.get('/api/bootstrap').json['cases']
        self.assertEqual({c['method'] for c in cases}, {'Administrativo', 'Extrajudicial'})
        response = self.client.get('/processos?method=Administrativo').data.decode()
        self.assertIn('1 cadastros encontrados', response)
        self.assertEqual(self.client.patch('/api/cases/' + cases[0]['id'], json={'stage': 'Teste', 'category': 'Perícia', 'author': 'Equipe'}).status_code, 400)

    def test_revision_preview_empty_values_immutability_and_stale_confirmation(self):
        initial = self.bundle()
        self.apply(initial)
        records = self.store.all('records')
        revision = self.bundle([{'G': 'Trecho 1', 'AU': 'Nova providência', 'H': None}, {'G': 'Trecho 2', 'P': 80}], 'revision.xlsx')
        plan = reg.plan_import(self.store, revision)
        self.assertEqual(plan['counts'], {'updated': 1, 'unchanged': 1})
        self.assertEqual([c['column'] for c in plan['items'][0]['changes']], ['AU'])
        reg.update_record(self.store, records[0]['id'], {'group': 'ATUALIZACAO_MVGA', 'version': '1', 'author': 'Equipe', 'reason': 'Correção documentada', 'BD': 'Equipe nova'})
        with self.assertRaises(Problem) as error:
            reg.apply_import(self.store, revision, plan['fingerprint'], 'Equipe')
        self.assertEqual(error.exception.status, 409)
        self.apply(revision)
        changed = self.store.get('records', records[0]['id'])
        self.assertEqual(changed['values']['H'], 'Pessoa fictícia')
        self.assertEqual(changed['values']['AU'], 'Nova providência')
        self.assertEqual(len(changed['sourceIds']), 2)
        self.assertTrue(reg.plan_import(self.store, initial)['alreadyImported'])
        with self.assertRaises(Problem):
            self.apply(initial)
        immutable = self.bundle([{'G': 'Trecho 1', 'E': 'TRINDADE'}], 'immutable.xlsx')
        self.assertEqual(reg.plan_import(self.store, immutable)['counts'], {'blocked': 1})

    def test_edit_governance_versioning_and_record_isolation(self):
        self.apply(self.bundle())
        first, other = self.store.all('records')
        base = {'group': 'FASE_PERICIAL', 'version': '1', 'author': 'Equipe', 'reason': 'Laudo recebido'}
        changed = reg.update_record(self.store, first['id'], {**base, 'AE': 'LAUDO APRESENTADO', 'AF': '', 'AB': '0,00'})
        self.assertEqual(changed['version'], 2)
        self.assertEqual(changed['values']['AB'], '0.00')
        self.assertEqual(self.store.get('records', other['id']), other)
        events = self.store.all('events', first['caseId'])
        self.assertTrue(any(change['before'] == 'PERITO NOMEADO' and change['after'] == 'LAUDO APRESENTADO'
                            for e in events for change in e.get('changes', [])))
        for bad in [
            {**base, 'AE': 'LAUDO PENDENTE'},  # stale version
            {**base, 'version': '2', 'AE': 'arbitrary invalid option'},
            {**base, 'version': '2', 'AO': '2026-01-01'},  # another group
            {**base, 'version': '2', 'AB': '-100'},
            {**base, 'version': '2', 'group': 'CADASTRO', 'C': 'new key'},
            {**base, 'version': '2', 'group': 'CADASTRO', 'D': '1000'},
            {**base, 'version': '2', 'group': 'CADASTRO', 'E': 'TRINDADE'},
            {**base, 'version': '2', 'group': 'ATE_IMISSAO_DE_POSSE', 'W': '10'},
            {**base, 'version': '2', 'group': 'ATE_IMISSAO_DE_POSSE', 'AP': 'LIBERADO'},
            {**base, 'version': '2', 'group': 'ATE_IMISSAO_DE_POSSE', 'AO': '2026-02-31'},
        ]:
            with self.subTest(bad=bad), self.assertRaises(Problem):
                reg.update_record(self.store, first['id'], bad)
        self.assertEqual(self.store.get('records', first['id']), changed)
        # Formal number correction retains the persistent key and does not update a sibling.
        corrected = reg.update_record(self.store, first['id'], {**base, 'version': '2', 'group': 'CADASTRO', 'L': '7654321-89.2025.8.09.0001'})
        self.assertEqual(corrected['key'], first['key'])
        self.assertEqual(self.store.get('records', other['id'])['values']['L'], CNJ)

    def test_import_rolls_back_all_changes_if_storage_fails(self):
        bundle = self.bundle()
        plan = reg.plan_import(self.store, bundle)
        original_put = self.store.put
        def fail(table, data, case_id=None):
            if table == 'events':
                raise RuntimeError('Simulated storage failure')
            return original_put(table, data, case_id)
        with patch.object(self.store, 'put', side_effect=fail), self.assertRaises(RuntimeError):
            reg.apply_import(self.store, bundle, plan['fingerprint'], 'Equipe')
        for table in ('cases', 'records', 'sources', 'events', 'meta'):
            self.assertEqual(self.store.all(table), [])

    def test_browser_forms_preview_confirmation_filters_and_escaping(self):
        path = workbook(self.root / 'control.xlsx', [{'F': '<script>alert(1)</script>', 'G': 'Área 1'}, {'G': 'Área 2'}])
        self.client.get('/importar')
        with self.client.session_transaction() as session:
            csrf = session['csrf']
        self.assertEqual(self.client.post('/importar', data={'workbook': (io.BytesIO(path.read_bytes()), 'control.xlsx')}).status_code, 403)
        response = self.client.post('/importar', data={'_csrf': csrf, 'workbook': (io.BytesIO(path.read_bytes()), 'control.xlsx')})
        self.assertEqual(response.status_code, 303)
        url = response.location
        page = self.client.get(url)
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'&lt;script&gt;', page.data)
        self.assertEqual(self.store.all('records'), [])
        self.assertEqual(self.app.test_client().get(url).status_code, 404)
        self.assertEqual(self.client.get(url + '?status=new&page=999').status_code, 200)
        result = self.client.post(url + '/confirmar', data={'_csrf': csrf, 'author': 'Equipe'})
        self.assertEqual(result.status_code, 303)
        records = self.store.all('records')
        self.assertEqual(len(records), 2)
        self.assertEqual(self.client.get(url).status_code, 404)
        for route in ['/fontes', '/', '/imoveis', '/imoveis?mandate=yes', '/imoveis?pending=1', '/processos', '/pericias', '/processos/' + records[0]['caseId']]:
            with self.subTest(route=route):
                self.assertEqual(self.client.get(route).status_code, 200)
        for group in reg.GROUPS:
            self.assertEqual(self.client.get(f"/imoveis/{records[0]['id']}?group={group}").status_code, 200)
        result = self.client.post('/imoveis/' + records[0]['id'], data={'_csrf': csrf, 'group': 'CADASTRO', 'version': '1', 'author': 'Equipe', 'reason': 'Peça técnica', 'N': 'REALIZAÇÃO DE PERÍCIA'})
        self.assertEqual(result.status_code, 303)
        self.assertIn(b'REALIZA', self.client.get('/imoveis?stage=REALIZA%C3%87%C3%83O+DE+PER%C3%8DCIA').data)


if __name__ == '__main__':
    unittest.main()
